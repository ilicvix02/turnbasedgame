"""Shared editing rights for TradeSheets workbooks.

Input cells (blue text, or yellow key-setting fill) stay editable; formula and label cells are locked so
customers can't overwrite a formula by accident. No password: Review → Unprotect Sheet turns it off.
"""
from openpyxl.styles import Protection

BLUE = "000000FF"
YELLOW = "00FFFF00"


def apply_editing_rights(wb, skip=("Start Here",)):
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for c in row:
                color = c.font.color.rgb if c.font and c.font.color is not None else None
                fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
                if color == BLUE or fill == YELLOW:
                    c.protection = Protection(locked=False)
        p = ws.protection
        p.sheet = True
        # Allowed even while protected: resizing, formatting, sorting/filtering, selecting.
        p.formatCells = p.formatColumns = p.formatRows = False
        p.sort = p.autoFilter = False
        p.selectLockedCells = p.selectUnlockedCells = False
