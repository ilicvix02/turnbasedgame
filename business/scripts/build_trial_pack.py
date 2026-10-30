"""Build the TradeSheets free trial pack (3 premium tools) for prospects who reply "I want it".

Usage:  python business/scripts/build_trial_pack.py [out_dir]
Writes 1-Job-Profit-Snapshot.xlsx, 2-90-Day-Cash-Flow.xlsx, 3-Truck-Stock-Unbilled-Parts.xlsx.
"""
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from protect import apply_editing_rights

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=16, color="0B2234")
NOTE = Font(name=FONT, italic=True, color="666666")
HEAD_FILL = PatternFill("solid", fgColor="0B2234")
GOLD_FILL = PatternFill("solid", fgColor="FFF2C2")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
AMBER_FILL = PatternFill("solid", fgColor="FFF3CD")
OK_FILL = PatternFill("solid", fgColor="D4EDDA")
THIN = Side(style="thin", color="C8D3DC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"
FOOTER = "TradeSheets free trial pack · Blue text = your inputs · Black = locked formulas (Review > Unprotect Sheet to change) · Full custom setup: reply to the email you got this from."


def brand(ws, title, subtitle, width_cols):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = subtitle
    ws["A2"].font = NOTE
    ws.sheet_view.showGridLines = False
    for col, w in width_cols.items():
        ws.column_dimensions[col].width = w


def header(ws, row, labels):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def cell(ws, ref, value, font=BLACK, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font, c.border = value, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def job_profit_snapshot(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Job Profit Snapshot"
    brand(ws, "Job Profit Snapshot", "Was this job worth it? Fill in the Bid column before you quote, and the Actual column when the job closes.",
          {"A": 38, "B": 16, "C": 16, "D": 16, "E": 44})
    header(ws, 4, ["", "Bid (estimate)", "Actual (job closed)", "Difference", "Notes"])
    inputs = [
        ("Job price charged to customer ($)", 6800, 6800, USD, "Include approved change orders in Actual."),
        ("Labor hours", 40, 46, "0.0", "All crew hours on the job, including drive time."),
        ("Loaded labor rate ($/hour)", 42, 42, USD2, "Wage + payroll taxes + workers' comp. Example value, use your own."),
        ("Materials & parts ($)", 2300, 2510, USD, ""),
        ("Subcontractors ($)", 0, 0, USD, ""),
        ("Equipment, rentals, permits, other ($)", 350, 420, USD, ""),
    ]
    for i, (label, bid, act, fmt, note) in enumerate(inputs):
        r = 5 + i
        cell(ws, f"A{r}", label)
        cell(ws, f"B{r}", bid, BLUE, fmt)
        cell(ws, f"C{r}", act, BLUE, fmt)
        cell(ws, f"D{r}", f"=C{r}-B{r}", BLACK, fmt)
        cell(ws, f"E{r}", note, NOTE)
    cell(ws, "A12", "Overhead as % of revenue", BOLD)
    cell(ws, "B12", 0.15, BLUE, PCT, INPUT_FILL)
    cell(ws, "E12", "Rent, insurance, trucks, phone, office ÷ yearly revenue. Example value, set your own.", NOTE)
    cell(ws, "A13", "Target net margin", BOLD)
    cell(ws, "B13", 0.15, BLUE, PCT, INPUT_FILL)
    cell(ws, "E13", "The profit you want left after overhead. Example value.", NOTE)

    header(ws, 15, ["Results", "Bid (estimate)", "Actual (job closed)", "Difference", "What it means"])
    results = [
        ("Labor cost ($)", "={c}6*{c}7", USD, "Hours × loaded rate."),
        ("Direct job cost ($)", "={c}16+SUM({c}8:{c}10)", USD, "Labor + materials + subs + other."),
        ("Gross profit ($)", "={c}5-{c}17", USD, "Before overhead."),
        ("Gross margin", '=IF({c}5=0,"",{c}18/{c}5)', PCT, ""),
        ("Overhead share ($)", "={c}5*$B$12", USD, "This job's fair share of running the business."),
        ("Net profit ($)", "={c}18-{c}20", USD, "What the job really earned."),
        ("Net margin", '=IF({c}5=0,"",{c}21/{c}5)', PCT, "Red if below your target."),
        ("Break-even price ($)", '=IF($B$12>=1,"",{c}17/(1-$B$12))', USD, "Charge less than this and the job loses money."),
        ("Price needed for target margin ($)", '=IF($B$12+$B$13>=1,"",{c}17/(1-$B$12-$B$13))', USD, "Quote at least this next time."),
    ]
    for i, (label, f, fmt, note) in enumerate(results):
        r = 16 + i
        cell(ws, f"A{r}", label, BOLD if label.startswith("Net") else BLACK)
        cell(ws, f"B{r}", f.format(c="B"), BLACK, fmt)
        cell(ws, f"C{r}", f.format(c="C"), BLACK, fmt)
        cell(ws, f"D{r}", f'=IF(OR(B{r}="",C{r}=""),"",C{r}-B{r})', BLACK, fmt)
        cell(ws, f"E{r}", note, NOTE)
    for col in "BC":
        ws.conditional_formatting.add(f"{col}22", FormulaRule(formula=[f'AND({col}22<>"",{col}22<$B$13)'], fill=RED_FILL))
        ws.conditional_formatting.add(f"{col}22", FormulaRule(formula=[f'AND({col}22<>"",{col}22>=$B$13)'], fill=OK_FILL))
    ws["A26"] = "Verdict"
    ws["A26"].font = BOLD
    ws["B26"] = '=IF(C21<0,"This job LOST money. Check labor hours first.",IF(C22<$B$13,"Profitable, but below your target margin.","Good job: at or above your target margin."))'
    ws["B26"].font = BOLD
    ws["A28"] = FOOTER
    ws["A28"].font = NOTE
    apply_editing_rights(wb)
    wb.save(out)


def cash_flow_90(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "90-Day Cash Flow"
    brand(ws, "90-Day Cash Flow Forecaster", "13 weeks of money in and money out. Red = the week your bank balance goes negative.",
          {"A": 12, "B": 13, "C": 16, "D": 13, "E": 13, "F": 15, "G": 15, "H": 13, "I": 14, "J": 16, "K": 22})
    cell(ws, "A4", "Starting bank balance ($)", BOLD)
    ws.merge_cells("A4:C4")
    cell(ws, "D4", 4000, BLUE, USD, INPUT_FILL)
    cell(ws, "A5", "First week starts on", BOLD)
    ws.merge_cells("A5:C5")
    cell(ws, "D5", date(2026, 11, 2), BLUE, DATE, INPUT_FILL)
    cell(ws, "A6", "Warn me below ($)", BOLD)
    ws.merge_cells("A6:C6")
    cell(ws, "D6", 3000, BLUE, USD, INPUT_FILL)
    cell(ws, "F4", "Example numbers are filled in (week 4 has a $3,200 truck repair). Replace the blue cells with your own.", NOTE)

    header(ws, 8, ["Week", "Week of", "Customer payments in ($)", "Other in ($)", "Total in ($)", "Payroll out ($)",
                   "Materials & suppliers ($)", "Overhead out ($)", "Other out ($)", "Ending balance ($)", "Status"])
    pay = [6500, 800, 8100, 2400, 7200, 1800, 9400, 3000, 6100, 2800, 8800, 2500, 7000]
    payroll = [4200, 0, 4200, 0, 4400, 0, 4400, 0, 4400, 0, 4600, 0, 4600]
    mats = [2600, 1800, 3100, 1500, 2900, 1200, 3600, 1400, 2400, 1700, 3300, 1500, 2800]
    ovh = [0, 0, 0, 1900, 0, 0, 0, 1900, 0, 0, 0, 1900, 0]
    for i in range(13):
        r = 9 + i
        cell(ws, f"A{r}", i + 1, BLACK, "0")
        cell(ws, f"B{r}", f"=$D$5+7*(A{r}-1)", BLACK, DATE)
        cell(ws, f"C{r}", pay[i], BLUE, USD)
        cell(ws, f"D{r}", 0, BLUE, USD)
        cell(ws, f"E{r}", f"=C{r}+D{r}", BLACK, USD)
        cell(ws, f"F{r}", payroll[i], BLUE, USD)
        cell(ws, f"G{r}", mats[i], BLUE, USD)
        cell(ws, f"H{r}", ovh[i], BLUE, USD)
        cell(ws, f"I{r}", 3200 if i == 3 else 0, BLUE, USD)
        prev = "$D$4" if i == 0 else f"J{r-1}"
        cell(ws, f"J{r}", f"={prev}+E{r}-SUM(F{r}:I{r})", BOLD, USD)
        cell(ws, f"K{r}", f'=IF(J{r}<0,"CASH SHORTFALL",IF(J{r}<$D$6,"Low: plan ahead","OK"))', BOLD)
    ws.conditional_formatting.add("J9:J21", CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL))
    ws.conditional_formatting.add("K9:K21", CellIsRule(operator="equal", formula=['"CASH SHORTFALL"'], fill=RED_FILL))
    ws.conditional_formatting.add("K9:K21", CellIsRule(operator="equal", formula=['"Low: plan ahead"'], fill=AMBER_FILL))
    ws.conditional_formatting.add("K9:K21", CellIsRule(operator="equal", formula=['"OK"'], fill=OK_FILL))
    cell(ws, "A23", "13-week totals", BOLD)
    ws.merge_cells("A23:B23")
    for col in "CDEFGHI":
        cell(ws, f"{col}23", f"=SUM({col}9:{col}21)", BOLD, USD)
    cell(ws, "A25", "Lowest balance", BOLD)
    ws.merge_cells("A25:B25")
    cell(ws, "C25", "=MIN(J9:J21)", BOLD, USD, GOLD_FILL)
    cell(ws, "D25", '=IF(C25<0,"Week "&INDEX(A9:A21,MATCH(C25,J9:J21,0))&": get a deposit in or move a supplier payment.","")', BOLD)
    ws["A27"] = FOOTER
    ws["A27"].font = NOTE
    ws.freeze_panes = "C9"
    apply_editing_rights(wb)
    wb.save(out)


def truck_stock(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts Log"
    brand(ws, "Truck Stock & Unbilled Parts Tracker", "Log every part pulled from a truck. Mark it billed when it's on an invoice. Anything unbilled is money you're owed.",
          {"A": 12, "B": 10, "C": 10, "D": 30, "E": 7, "F": 12, "G": 12, "H": 13, "I": 9, "J": 12, "K": 14})
    cell(ws, "A4", "Parts markup %", BOLD)
    ws.merge_cells("A4:C4")
    cell(ws, "D4", 0.35, BLUE, PCT, INPUT_FILL)
    cell(ws, "E4", "What you add on top of cost when you bill parts. Example value.", NOTE)
    header(ws, 6, ["Date", "Truck", "Job ID", "Part / description", "Qty", "Unit cost ($)", "Cost ($)", "Bill price ($)",
                   "Billed?", "Invoice #", "Unbilled ($)"])
    ex = [
        (date(2026, 11, 2), "T1", "J-201", '3/4" ball valve', 2, 18.40, "Y", "INV-311"),
        (date(2026, 11, 2), "T1", "J-201", "Braided supply lines", 4, 7.25, "N", ""),
        (date(2026, 11, 3), "T2", "J-202", "40 gal water heater kit", 1, 46.00, "Y", "INV-312"),
        (date(2026, 11, 3), "T2", "J-203", "PEX fittings (bag)", 1, 32.80, "N", ""),
        (date(2026, 11, 4), "T1", "J-204", "Wax ring + bolts", 2, 6.90, "N", ""),
        (date(2026, 11, 4), "T3", "J-205", "20A GFCI receptacle", 3, 21.50, "Y", "INV-314"),
    ]
    last = 306
    for r in range(7, last + 1):
        if r - 7 < len(ex):
            d, t, j, part, q, u, billed, inv = ex[r - 7]
            for col, v in zip("ABCDEFIJ", (d, t, j, part, q, u, billed, inv)):
                ws[f"{col}{r}"] = v
        ws[f"G{r}"] = f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'
        ws[f"H{r}"] = f'=IF(G{r}="","",G{r}*(1+$D$4))'
        ws[f"K{r}"] = f'=IF(AND(G{r}<>"",I{r}<>"Y"),H{r},"")'
        for col, fmt, font in [("A", DATE, BLUE), ("B", None, BLUE), ("C", None, BLUE), ("D", None, BLUE), ("E", "0", BLUE),
                               ("F", USD2, BLUE), ("G", USD2, BLACK), ("H", USD2, BLACK), ("I", None, BLUE), ("J", None, BLUE),
                               ("K", USD2, BOLD)]:
            c = ws[f"{col}{r}"]
            c.font, c.border = font, BOX
            if fmt:
                c.number_format = fmt
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I7:I{last}")
    ws.conditional_formatting.add(f"K7:K{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))
    ws.freeze_panes = "A7"

    sm = wb.create_sheet("Summary")
    brand(sm, "Unbilled Parts Summary", "Updates automatically from the Parts Log.", {"A": 30, "B": 18, "C": 18, "D": 18})
    rng = f"'Parts Log'!$K$7:$K${last}"
    cell(sm, "A4", "Total parts cost logged ($)", BOLD)
    cell(sm, "B4", f"=SUM('Parts Log'!$G$7:$G${last})", BOLD, USD2)
    cell(sm, "A5", "Unbilled parts at bill price ($)", BOLD)
    cell(sm, "B5", f"=SUM({rng})", BOLD, USD2, GOLD_FILL)
    cell(sm, "A6", "Unbilled line items", BOLD)
    cell(sm, "B6", f'=COUNTIFS(\'Parts Log\'!$I$7:$I${last},"<>Y",\'Parts Log\'!$G$7:$G${last},">0")', BOLD, "0")
    header(sm, 8, ["Truck", "Parts cost ($)", "Unbilled ($)", "% unbilled"])
    for i, t in enumerate(["T1", "T2", "T3", "T4", "T5"]):
        r = 9 + i
        cell(sm, f"A{r}", t, BLUE)
        cell(sm, f"B{r}", f"=SUMIFS('Parts Log'!$G$7:$G${last},'Parts Log'!$B$7:$B${last},A{r})", BLACK, USD2)
        cell(sm, f"C{r}", f"=SUMIFS({rng},'Parts Log'!$B$7:$B${last},A{r})", BLACK, USD2)
        cell(sm, f"D{r}", f'=IF(B{r}=0,"",C{r}/(B{r}*(1+\'Parts Log\'!$D$4)))', BLACK, PCT)
    cell(sm, "A15", "Type your own truck names in the blue cells above.", NOTE)
    sm["A17"] = FOOTER
    sm["A17"].font = NOTE
    apply_editing_rights(wb)
    wb.save(out)


if __name__ == "__main__":
    out_dir = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "email-assets" / "trial-pack")
    out_dir.mkdir(parents=True, exist_ok=True)
    job_profit_snapshot(out_dir / "1-Job-Profit-Snapshot.xlsx")
    cash_flow_90(out_dir / "2-90-Day-Cash-Flow.xlsx")
    truck_stock(out_dir / "3-Truck-Stock-Unbilled-Parts.xlsx")
    print("built", out_dir)
